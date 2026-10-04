from __future__ import annotations

import hashlib
import io
import re
import urllib.parse
import urllib.request
from dataclasses import asdict, dataclass
from datetime import date
from html.parser import HTMLParser

from openpyxl import load_workbook


JPX_LISTED_PAGE_URL = "https://www.jpx.co.jp/english/markets/statistics-equities/misc/01.html"


@dataclass(frozen=True)
class ListedIssue:
    code: str
    company_name: str
    market_segment: str
    snapshot_date: str

    @property
    def symbol(self) -> str:
        return f"{self.code}.T"

    def to_dict(self) -> dict:
        return asdict(self) | {"symbol": self.symbol}


class _PageParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []
        self.text_parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)

    def handle_data(self, data):
        self.text_parts.append(data)


def _normalize_code(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _snapshot_date_from_text(text: str) -> date:
    match = re.search(
        r"List of TSE-listed Issues\s*\(([A-Z][a-z]{2})\.\s*(\d{4})\)",
        " ".join(text.split()),
    )
    if not match:
        raise ValueError("JPX listed-issues snapshot month not found")
    month = {
        "Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4,
        "May": 5, "Jun": 6, "Jul": 7, "Aug": 8,
        "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12,
    }[match.group(1)]
    year = int(match.group(2))
    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)
    return date.fromordinal(next_month.toordinal() - 1)


def parse_listed_page(raw: bytes) -> tuple[str, date]:
    parser = _PageParser()
    parser.feed(raw.decode("utf-8", errors="replace"))
    xlsx = next(
        (
            urllib.parse.urljoin(JPX_LISTED_PAGE_URL, href)
            for href in parser.links
            if href.lower().endswith(".xlsx") and "data_e" in href.lower()
        ),
        None,
    )
    if not xlsx:
        raise ValueError("JPX listed-issues Excel link not found")
    snapshot = _snapshot_date_from_text(" ".join(parser.text_parts))
    return xlsx, snapshot


def _header_columns(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows[:50]):
        labels = {
            str(value or "").strip().lower(): col_index
            for col_index, value in enumerate(row)
            if str(value or "").strip()
        }
        code_col = next(
            (
                col
                for label, col in labels.items()
                if label == "code" or "issue code" in label or "security code" in label
            ),
            None,
        )
        name_col = next(
            (
                col
                for label, col in labels.items()
                if "issue name" in label or "company name" in label or label == "name"
            ),
            None,
        )
        market_col = next(
            (
                col
                for label, col in labels.items()
                if "market" in label or "product" in label
            ),
            None,
        )
        if code_col is not None and name_col is not None and market_col is not None:
            return row_index, {"code": code_col, "name": name_col, "market": market_col}
    return None


def _infer_columns(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    code_pattern = re.compile(r"^(?:[0-9]{4}|[0-9]{3}[A-Z])$")
    sample = rows[:800]
    col_count = max((len(row) for row in sample), default=0)
    if not col_count:
        raise ValueError("empty JPX listed-issues sheet")

    code_scores = [0] * col_count
    market_scores = [0] * col_count
    text_scores = [0] * col_count

    for row in sample:
        for col, value in enumerate(row):
            if value is None:
                continue
            normalized = _normalize_code(value)
            lower = normalized.lower()
            if code_pattern.match(normalized):
                code_scores[col] += 1
            if any(label in lower for label in ("prime", "standard", "growth")):
                market_scores[col] += 1
            if isinstance(value, str) and len(value.strip()) >= 2:
                text_scores[col] += 1

    code_col = max(range(col_count), key=lambda col: code_scores[col])
    market_col = max(range(col_count), key=lambda col: market_scores[col])
    if code_scores[code_col] < 20 or market_scores[market_col] < 20:
        raise ValueError("unable to infer JPX listed-issues columns")

    preferred = code_col + 1
    if preferred < col_count and preferred != market_col and text_scores[preferred] >= 20:
        name_col = preferred
    else:
        candidates = [
            col
            for col in range(col_count)
            if col not in {code_col, market_col}
        ]
        name_col = max(candidates, key=lambda col: text_scores[col])

    return -1, {"code": code_col, "name": name_col, "market": market_col}


def parse_listed_workbook(raw: bytes, *, snapshot_date: date) -> list[ListedIssue]:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    failures: list[str] = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        try:
            detected = _header_columns(rows)
            header_row, cols = detected if detected is not None else _infer_columns(rows)
        except ValueError as exc:
            failures.append(f"{ws.title}: {exc}")
            continue

        issues: list[ListedIssue] = []
        seen: set[str] = set()
        start_index = header_row + 1 if header_row >= 0 else 0

        for values in rows[start_index:]:
            def cell(index: int):
                return values[index] if index < len(values) else None

            code = _normalize_code(cell(cols["code"]))
            name = str(cell(cols["name"]) or "").strip()
            market = str(cell(cols["market"]) or "").strip()
            lower = market.lower()

            if not code or not name:
                continue
            if not any(label in lower for label in ("prime", "standard", "growth")):
                continue
            if "foreign" in lower or "pro market" in lower:
                continue
            if code in seen:
                continue

            seen.add(code)
            issues.append(
                ListedIssue(
                    code=code,
                    company_name=name,
                    market_segment=market,
                    snapshot_date=snapshot_date.isoformat(),
                )
            )

        if issues:
            return issues
        failures.append(f"{ws.title}: zero domestic stock rows")

    raise ValueError(
        "JPX listed-issues workbook could not be parsed; "
        + "; ".join(failures[:8])
    )


def fetch_current_listed_issues(
    *,
    timeout_seconds: int = 30,
) -> tuple[list[ListedIssue], dict]:
    req = urllib.request.Request(
        JPX_LISTED_PAGE_URL,
        headers={"User-Agent": "Mozilla/5.0 stock-warashibe point-in-time/0.7"},
    )
    with urllib.request.urlopen(req, timeout=timeout_seconds) as response:
        page_raw = response.read()

    xlsx_url, snapshot_date = parse_listed_page(page_raw)
    req = urllib.request.Request(
        xlsx_url,
        headers={"User-Agent": "Mozilla/5.0 stock-warashibe point-in-time/0.7"},
    )
    with urllib.request.urlopen(req, timeout=timeout_seconds) as response:
        workbook_raw = response.read()

    issues = parse_listed_workbook(workbook_raw, snapshot_date=snapshot_date)
    return issues, {
        "page_url": JPX_LISTED_PAGE_URL,
        "xlsx_url": xlsx_url,
        "snapshot_date": snapshot_date.isoformat(),
        "page_sha256": hashlib.sha256(page_raw).hexdigest(),
        "workbook_sha256": hashlib.sha256(workbook_raw).hexdigest(),
        "record_count": len(issues),
    }
